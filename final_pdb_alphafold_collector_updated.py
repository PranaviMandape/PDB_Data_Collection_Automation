# ============================================================
# PDB + AlphaFold Protein Structure Data Collector
# ============================================================
# USER INPUT: change only protein_name.
# Output: Protein_Structure_Data_<protein>.xlsx
# Notes / Errors column is intentionally removed.
# PDB Chains is placed immediately after Structure Type.
# Modeled Residue Count is the sum of chain-level modeled_residue_count values.
# ============================================================

# %% 1. Imports and setup
import sys, subprocess, io, gzip, re, time, warnings
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from collections import defaultdict
from datetime import datetime, timezone
from functools import lru_cache


def ensure_package(import_name, package_name=None):
    try:
        __import__(import_name)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", package_name or import_name])

for mod, pkg in [("requests", "requests"), ("pandas", "pandas"), ("openpyxl", "openpyxl"),
                 ("pdfplumber", "pdfplumber"), ("Bio", "biopython")]:
    ensure_package(mod, pkg)

import requests
import pandas as pd
import pdfplumber
import xml.etree.ElementTree as ET
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from Bio.PDB.MMCIF2Dict import MMCIF2Dict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# %% 2. User input and settings
protein_name = "ITK"                 # <-- CHANGE ONLY THIS
HUMAN_TAXONOMY_ID = 9606
REQUEST_TIMEOUT = 30
MAX_RETRIES = 3
RETRY_BACKOFF = 1.5
RATE_LIMIT_DELAY = 0.15
OUTPUT_FILENAME = f"Protein_Structure_Data_{protein_name}.xlsx"

print(f"Protein: {protein_name}")
print(f"Output:  {OUTPUT_FILENAME}")


# %% 3. API endpoints and HTTP helpers
UNIPROT_SEARCH_URL = "https://rest.uniprot.org/uniprotkb/search"
UNIPROT_ENTRY_URL = "https://rest.uniprot.org/uniprotkb/{accession}.json"
UNIPROT_BROWSER_URL = "https://www.uniprot.org/uniprotkb/{accession}"

RCSB_SEARCH_URL = "https://search.rcsb.org/rcsbsearch/v2/query"
RCSB_ENTRY_URL = "https://data.rcsb.org/rest/v1/core/entry/{pdb_id}"
RCSB_POLYMER_URL = "https://data.rcsb.org/rest/v1/core/polymer_entity/{pdb_id}/{entity_id}"
RCSB_POLYMER_INSTANCE_URL = "https://data.rcsb.org/rest/v1/core/polymer_entity_instance/{pdb_id}/{asym_id}"
RCSB_NONPOLY_URL = "https://data.rcsb.org/rest/v1/core/nonpolymer_entity/{pdb_id}/{entity_id}"
RCSB_ASSEMBLY_URL = "https://data.rcsb.org/rest/v1/core/assembly/{pdb_id}/{assembly_id}"
MMCIF_URL = "https://files.rcsb.org/download/{pdb_id}.cif"

VALIDATION_XML_URL = "https://files.rcsb.org/pub/pdb/validation_reports/{mid}/{pid}/{pid}_validation.xml.gz"
VALIDATION_PDF_URL = "https://files.rcsb.org/pub/pdb/validation_reports/{mid}/{pid}/{pid}_full_validation.pdf.gz"

ALPHAFOLD_API_URL = "https://alphafold.com/api/prediction/{accession}"
ALPHAFOLD_ENTRY_URL = "https://alphafold.com/entry/{identifier}"
THREED_BEACONS_URL = "https://www.ebi.ac.uk/pdbe/pdbe-kb/3dbeacons/api/uniprot/summary/{accession}.json"


def make_session():
    s = requests.Session()
    retry = Retry(total=MAX_RETRIES, connect=MAX_RETRIES, read=MAX_RETRIES,
                  backoff_factor=RETRY_BACKOFF,
                  status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=frozenset({"GET", "POST"}),
                  raise_on_status=False)
    adapter = HTTPAdapter(max_retries=retry)
    s.mount("https://", adapter)
    s.headers.update({"User-Agent": "ProteinPDBDataCollector/2.0"})
    return s

SESSION = make_session()


def request_json(url, params=None, timeout=REQUEST_TIMEOUT):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.get(url, params=params, timeout=timeout)
        r.raise_for_status()
        return r.json(), None
    except (requests.RequestException, ValueError) as e:
        return None, str(e)


def request_bytes(url, timeout=REQUEST_TIMEOUT):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.get(url, timeout=timeout)
        r.raise_for_status()
        return r.content, None
    except requests.RequestException as e:
        return None, str(e)


def rcsb_search(query):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.post(RCSB_SEARCH_URL, json=query, timeout=REQUEST_TIMEOUT,
                         headers={"Content-Type": "application/json"})
        r.raise_for_status()
        return r.json(), None
    except (requests.RequestException, ValueError) as e:
        return None, str(e)

print("API/session setup complete.")


# %% 4. UniProt human entry resolution
def search_uniprot(gene, reviewed=True):
    q = [f"gene_exact:{gene}", f"organism_id:{HUMAN_TAXONOMY_ID}"]
    if reviewed:
        q.append("reviewed:true")
    data, err = request_json(UNIPROT_SEARCH_URL, {
        "query": " AND ".join(q), "format": "json", "size": 25,
        "fields": "accession,id,protein_name,organism_name,gene_names,reviewed,organism_id"
    })
    return (data or {}).get("results", []) if not err else []


def resolve_uniprot(gene):
    matches = search_uniprot(gene, True) or search_uniprot(gene, False)
    human = [x for x in matches if str(x.get("organism", {}).get("taxonId", x.get("organism_id", ""))) == str(HUMAN_TAXONOMY_ID)]
    if not human:
        raise RuntimeError(f"No human UniProt entry found for exact gene '{gene}'.")
    if len(human) > 1:
        ids = [x.get("primaryAccession", "?") for x in human]
        raise RuntimeError(f"Multiple human UniProt entries found for '{gene}': {ids}")
    x = human[0]
    acc = x.get("primaryAccession")
    if not acc:
        raise RuntimeError("UniProt response did not contain an accession.")
    return acc, HUMAN_TAXONOMY_ID

UNIPROT_ACCESSION, ORGANISM_ID = resolve_uniprot(protein_name)
PROTEIN_NAME_COLUMN_VALUE = f"{protein_name} ({UNIPROT_ACCESSION})"
print(f"UniProt: {UNIPROT_ACCESSION} | Taxonomy: {ORGANISM_ID}")


# %% 5. Exact-value helpers
def raw_value(value):
    """Keep source values as text; do not round or invent digits."""
    if value is None:
        return "N/A"
    value = str(value).strip()
    return "N/A" if value.lower() in {"", "n/a", "na", ".", "?", "none"} else value


def decimal_from_text(value):
    if value is None or str(value).strip().lower() in {"", "n/a", "na", ".", "?", "none"}:
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        return None


def exact_difference(a, b):
    """Calculate a-b using Decimal without floating-point rounding."""
    x, y = decimal_from_text(a), decimal_from_text(b)
    if x is None or y is None:
        return "N/A"
    return format(x - y, "f")


# %% 6. RCSB X-ray structure search
# Uses POST because RCSB Search API expects the JSON query as the request body.
def get_pdb_ids(accession):
    query = {
        "query": {
            "type": "group", "logical_operator": "and", "nodes": [
                {"type": "terminal", "service": "text", "parameters": {
                    "attribute": "rcsb_polymer_entity_container_identifiers.reference_sequence_identifiers.database_accession",
                    "operator": "exact_match", "value": accession}},
                {"type": "terminal", "service": "text", "parameters": {
                    "attribute": "exptl.method", "operator": "exact_match", "value": "X-RAY DIFFRACTION"}}
            ]
        },
        "return_type": "entry",
        "request_options": {"paginate": {"start": 0, "rows": 10000}}
    }
    data, err = rcsb_search(query)
    if err:
        raise RuntimeError(f"RCSB Search API failed: {err}")
    ids = []
    for item in (data or {}).get("result_set", []):
        pid = str(item.get("identifier", "")).upper().strip()
        if pid and pid not in ids:
            ids.append(pid)
    return [{"pdb_id": x} for x in ids]

pdb_id_records = get_pdb_ids(UNIPROT_ACCESSION)
print(f"X-ray PDB structures found: {len(pdb_id_records)}")


# %% 7. RCSB entry, method, resolution, residue counts and chains
def get_entry(pdb_id):
    return request_json(RCSB_ENTRY_URL.format(pdb_id=pdb_id))[0]


def get_method(entry):
    methods = [x.get("method", "").strip() for x in entry.get("exptl", []) if x.get("method")]
    return ", ".join(dict.fromkeys(methods)) or "N/A"


def get_resolution(entry):
    vals = (entry.get("rcsb_entry_info") or {}).get("resolution_combined") or []
    vals = vals if isinstance(vals, list) else [vals]
    for value in vals:
        value = raw_value(value)
        if value != "N/A":
            return value
    return "N/A"


def get_residue_counts(entry, pdb_id):
    """
    Deposited Residue Count = RCSB entry deposited polymer monomer count.
    Modeled Residue Count = SUM of modeled_residue_count for every deposited
    polymer chain (polymer entity instance).
    Difference = Deposited - Modeled.
    """
    deposited_raw = (entry.get("rcsb_entry_info") or {}).get(
        "deposited_polymer_monomer_count"
    )

    try:
        deposited = int(deposited_raw) if deposited_raw is not None else None
    except (TypeError, ValueError):
        deposited = None

    if deposited is None:
        return None, None, None

    polymer_entity_ids = (
        entry.get("rcsb_entry_container_identifiers") or {}
    ).get("polymer_entity_ids") or []

    modeled_total = 0
    found_modeled_value = False
    seen_asym_ids = set()

    for entity_id in polymer_entity_ids:
        entity, err = request_json(
            RCSB_POLYMER_URL.format(
                pdb_id=pdb_id.upper(), entity_id=entity_id
            )
        )
        if err or not entity:
            continue

        asym_ids = (
            entity.get("rcsb_polymer_entity_container_identifiers") or {}
        ).get("asym_ids") or []
        if isinstance(asym_ids, str):
            asym_ids = [asym_ids]

        for asym_id in asym_ids:
            asym_id = str(asym_id).strip()
            if not asym_id or asym_id in seen_asym_ids:
                continue
            seen_asym_ids.add(asym_id)

            instance, err = request_json(
                RCSB_POLYMER_INSTANCE_URL.format(
                    pdb_id=pdb_id.upper(), asym_id=asym_id
                )
            )
            if err or not instance:
                continue

            modeled = (
                instance.get("rcsb_polymer_instance_info") or {}
            ).get("modeled_residue_count")

            if modeled is None:
                continue

            try:
                modeled_total += int(modeled)
                found_modeled_value = True
            except (TypeError, ValueError):
                continue

    modeled = modeled_total if found_modeled_value else None
    difference = deposited - modeled if modeled is not None else None
    return deposited, modeled, difference


def get_pdb_chains(entry, pdb_id):
    ids = (entry.get("rcsb_entry_container_identifiers") or {}).get("polymer_entity_ids") or []
    chains = []
    for entity_id in ids:
        data, err = request_json(RCSB_POLYMER_URL.format(pdb_id=pdb_id.upper(), entity_id=entity_id))
        if err or not data:
            continue
        asym = (data.get("rcsb_polymer_entity_container_identifiers") or {}).get("asym_ids") or []
        asym = [asym] if isinstance(asym, str) else asym
        for chain in asym:
            chain = str(chain).strip()
            if chain and chain not in chains:
                chains.append(chain)
    return ", ".join(chains) if chains else "N/A"


# %% 8. Depositor R-values from mmCIF
def get_depositor_r_values(pdb_id):
    raw, err = request_bytes(MMCIF_URL.format(pdb_id=pdb_id.upper()))
    if err or raw is None:
        return None, None
    try:
        cif = MMCIF2Dict(io.StringIO(raw.decode("utf-8", errors="replace")))
        def first(key):
            v = cif.get(key)
            if isinstance(v, list):
                v = next((x for x in v if str(x).strip() not in {"", ".", "?"}), None)
            return None if v is None or str(v).strip() in {"", ".", "?"} else str(v).strip()
        return first("_refine.ls_R_factor_R_free"), first("_refine.ls_R_factor_R_work")
    except Exception:
        return None, None


# %% 9. wwPDB validation XML + Average B-factor
AVERAGE_B_PATTERN = re.compile(r"Average\s*B,\s*all\s*atoms\s*\([^)]*\)\s*(-?\d+\.\d+)", re.I)


def get_validation(pdb_id):
    pid = pdb_id.lower(); mid = pid[1:3]
    raw, err = request_bytes(VALIDATION_XML_URL.format(mid=mid, pid=pid))
    if err or raw is None:
        return {}, None
    try:
        root = ET.fromstring(gzip.decompress(raw))
        node = root.find("Entry")
        return (dict(node.attrib), None) if node is not None else ({}, "no Entry element")
    except Exception as e:
        return {}, str(e)


def get_average_b(pdb_id):
    pid = pdb_id.lower(); mid = pid[1:3]
    raw, err = request_bytes(VALIDATION_PDF_URL.format(mid=mid, pid=pid))
    if err or raw is None:
        return "N/A"
    try:
        with pdfplumber.open(io.BytesIO(gzip.decompress(raw))) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                m = AVERAGE_B_PATTERN.search(text)
                if m:
                    return m.group(1)
                for line in text.splitlines():
                    if "Average B" in line and "atom" in line.lower():
                        nums = re.findall(r"-?\d+\.\d+", line)
                        if nums:
                            return nums[-1]
    except Exception:
        pass
    return "N/A"


# %% 10. Ligand and binding-activity extraction
COMMON_NON_LIGAND_COMPONENTS = {
    "HOH","H2O","WAT","NA","CL","K","MG","CA","ZN","MN","CD","CO","NI","CU","FE","FE2",
    "BR","IOD","CS","LI","SR","BA","RB","AL","AG","AU","HG","SO4","PO4","GOL","EDO","PEG",
    "PG4","MPD","ACT","TRS","IPA","DMS","BME","MES","HEPES","CIT","FMT","ACY","1PE","PGE","P6G",
    "EPE","IMD","NH4","OXL","UNX","UNL","DTT","TAM","BOG","CO3","NO3","AZI"
}


def get_ligand_info(entry, pdb_id):
    ids = (entry.get("rcsb_entry_container_identifiers") or {}).get("non_polymer_entity_ids") or []
    raw = []
    for entity_id in ids:
        data, err = request_json(RCSB_NONPOLY_URL.format(pdb_id=pdb_id, entity_id=entity_id))
        if err or not data:
            continue
        info = data.get("pdbx_entity_nonpoly") or {}
        comp = str(info.get("comp_id", "")).strip()
        if comp:
            raw.append({"comp_id": comp, "name": info.get("name") or comp})
    relevant = [x for x in raw if x["comp_id"].upper() not in COMMON_NON_LIGAND_COMPONENTS]
    return relevant, raw


def get_ligand_activity(entry, ligand_ids):
    affinities = entry.get("rcsb_binding_affinity") or []
    if not affinities or not ligand_ids:
        return "N/A"
    grouped = defaultdict(lambda: defaultdict(list))
    for a in affinities:
        comp = a.get("comp_id", "")
        if comp in ligand_ids:
            grouped[comp][a.get("type", "?")].append((a.get("value"), a.get("unit", "")))
    parts = []
    for comp, types in grouped.items():
        assay_parts = []
        for assay, vals in types.items():
            nums = [v for v, _ in vals if v is not None]
            unit = vals[0][1] if vals else ""
            if len(nums) == 1:
                assay_parts.append(f"{assay}: {nums[0]} ({unit}) from {len(vals)} assay(s)")
            elif nums:
                assay_parts.append(f"{assay}: min: {min(nums)}, max: {max(nums)} ({unit}) from {len(vals)} assay(s)")
        if assay_parts:
            parts.append(f"[{comp}] " + "; ".join(assay_parts))
    return " | ".join(parts) if parts else "N/A"


# %% 11. AlphaFold helpers
def get_alphafold_predictions(accession):
    data, err = request_json(ALPHAFOLD_API_URL.format(accession=accession))
    if err or data is None:
        return []
    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict)]
    if isinstance(data, dict):
        for key in ("predictions", "results", "data", "entries"):
            if isinstance(data.get(key), list):
                return [x for x in data[key] if isinstance(x, dict)]
        return [data]
    return []


def recursive_records(data):
    out = []
    if isinstance(data, dict):
        if any(k in data for k in ("model_identifier","modelIdentifier","model_id","modelId","modelEntityId","entryId","identifier")):
            out.append(data)
        for v in data.values(): out.extend(recursive_records(v))
    elif isinstance(data, list):
        for x in data: out.extend(recursive_records(x))
    return out


def model_id(record):
    for k in ("model_identifier","modelIdentifier","model_id","modelId","modelEntityId","entryId","identifier"):
        if record.get(k): return str(record[k]).strip()
    return ""


def af_positions(record):
    start = record.get("uniprot_start", record.get("sequenceStart", record.get("uniprotStart")))
    end = record.get("uniprot_end", record.get("sequenceEnd", record.get("uniprotEnd")))
    return f"{start}-{end}" if start is not None and end is not None else ""


def af_oligomeric_state(record):
    for k in ("oligomeric_state","oligomericState","oligomer_state","oligomerState"):
        v = record.get(k)
        if isinstance(v, dict):
            for sub in ("name","state","value","description"):
                if v.get(sub): return str(v[sub]).strip()
        elif isinstance(v, list):
            vals = [str(x).strip() for x in v if str(x).strip()]
            if vals: return ", ".join(vals)
        elif v is not None and str(v).strip():
            return str(v).strip()
    return ""


def structure_urls(data):
    cif, pdb = [], []
    def add(lst, v):
        if isinstance(v, str) and v.startswith(("http://", "https://")) and v not in lst: lst.append(v)
    def walk(o):
        if isinstance(o, dict):
            for k, v in o.items():
                kl, vl = str(k).lower(), str(v).lower() if isinstance(v, str) else ""
                if isinstance(v, str):
                    if "cif" in kl or vl.endswith(".cif") or ".cif?" in vl: add(cif, v)
                    elif "pdb" in kl or vl.endswith(".pdb") or ".pdb?" in vl: add(pdb, v)
                walk(v)
        elif isinstance(o, list):
            for x in o: walk(x)
    walk(data)
    return cif, pdb


def af_chain_from_url(url):
    if not url: return ""
    raw, err = request_bytes(url)
    if err or raw is None: return ""
    try:
        low = url.lower()
        if ".pdb" in low or raw.lstrip().startswith((b"HEADER", b"ATOM")):
            chains = []
            for line in raw.decode("utf-8", errors="replace").splitlines():
                if line.startswith(("ATOM", "HETATM")) and len(line) >= 22:
                    c = line[21].strip()
                    if c and c not in chains: chains.append(c)
            return ", ".join(chains)
        cif = MMCIF2Dict(io.StringIO(raw.decode("utf-8", errors="replace")))
        for key in ("_atom_site.auth_asym_id", "_atom_site.label_asym_id"):
            vals = cif.get(key)
            if vals:
                vals = vals if isinstance(vals, list) else [vals]
                chains = list(dict.fromkeys(str(x).strip() for x in vals if str(x).strip() not in {".", "?"}))
                if chains: return ", ".join(chains)
    except Exception:
        pass
    return ""


def extract_alphafold_rows(accession, protein_display):
    api_data = get_alphafold_predictions(accession)
    beacon, _ = request_json(THREED_BEACONS_URL.format(accession=accession))
    candidates = recursive_records(beacon) if beacon else []
    candidates = [x for x in candidates if model_id(x) or any("alphafold" in str(x.get(k, "")).lower() for k in ("provider","provider_name","providerName"))]
    if not candidates: candidates = api_data
    rows, seen = [], set()
    for rec in candidates:
        ident = model_id(rec)
        if not ident or ident.upper() in seen: continue
        seen.add(ident.upper())
        cif_urls, pdb_urls = structure_urls(rec)
        if not cif_urls and not pdb_urls:
            for src in api_data:
                if model_id(src).upper() == ident.upper():
                    cif_urls, pdb_urls = structure_urls(src); break
        url = cif_urls[0] if cif_urls else (pdb_urls[0] if pdb_urls else "")
        rows.append({"Identifier": ident,
                     "Chain": af_chain_from_url(url), "Positions": af_positions(rec),
                     "Oligomeric State": af_oligomeric_state(rec)})
    return rows


# %% 12. Column definitions (Notes / Errors removed)
FINAL_COLUMNS = [
    "Protein Name", "PDB ID", "PDB Link", "Method", "Resolution (Å)",
    "Deposited Residue Count", "Modeled Residue Count", "Residue Count Difference",
    "R-Value Free", "R-Value Work", "R-Value Difference", "R-free",
    "Clashscore", "Ramachandran Outliers (%)", "Sidechain Outliers (%)", "RSRZ Outliers (%)",
    "Average B, all atoms (Å²)", "Ligand ID", "Ligand Name", "Structure Type", "PDB Chains", "Ligand Activity",
    "All HET Components (raw)"
]

print(f"PDB_Data columns: {len(FINAL_COLUMNS)}")


# %% 13. Main PDB processing
rows = []
failed = []
for i, rec in enumerate(pdb_id_records, 1):
    pid = rec["pdb_id"]
    print(f"[{i}/{len(pdb_id_records)}] {pid} ...", end=" ")
    row = {c: "N/A" for c in FINAL_COLUMNS}
    row["Protein Name"] = PROTEIN_NAME_COLUMN_VALUE
    row["PDB ID"] = pid
    row["PDB Link"] = pid
    try:
        entry, err = request_json(RCSB_ENTRY_URL.format(pdb_id=pid))
        if err or not entry:
            failed.append((pid, f"RCSB entry: {err}")); rows.append(row); print("FAILED"); continue

        row["Method"] = get_method(entry)
        row["Resolution (Å)"] = get_resolution(entry)
        dep, mod, diff = get_residue_counts(entry, pid)
        row["Deposited Residue Count"] = dep if dep is not None else "N/A"
        row["Modeled Residue Count"] = mod if mod is not None else "N/A"
        row["Residue Count Difference"] = diff if diff is not None else "N/A"

        chains = get_pdb_chains(entry, pid)
        validation, _ = get_validation(pid)
        rfree, rwork = get_depositor_r_values(pid)
        row["R-Value Free"] = raw_value(rfree)
        row["R-Value Work"] = raw_value(rwork)
        row["R-Value Difference"] = exact_difference(row["R-Value Free"], row["R-Value Work"])
        row["R-free"] = raw_value(validation.get("DCC_Rfree"))
        row["Clashscore"] = raw_value(validation.get("clashscore"))
        row["Ramachandran Outliers (%)"] = raw_value(validation.get("percent-rama-outliers"))
        row["Sidechain Outliers (%)"] = raw_value(validation.get("percent-rota-outliers"))
        row["RSRZ Outliers (%)"] = raw_value(validation.get("percent-RSRZ-outliers"))
        row["Average B, all atoms (Å²)"] = get_average_b(pid)

        ligands, raw = get_ligand_info(entry, pid)
        row["All HET Components (raw)"] = ", ".join(x["comp_id"] for x in raw) if raw else "N/A"
        if ligands:
            row["Ligand ID"] = ", ".join(x["comp_id"] for x in ligands)
            row["Ligand Name"] = ", ".join(x["name"] for x in ligands)
            row["Structure Type"] = "Complex"
            row["Ligand Activity"] = get_ligand_activity(entry, {x["comp_id"] for x in ligands})
        else:
            row["Structure Type"] = "Apo"
        # PDB Chains intentionally placed next to Structure Type in the final Excel order.
        row["PDB Chains"] = chains
        rows.append(row)
        print("done")
    except Exception as e:
        failed.append((pid, str(e))); rows.append(row); print(f"FAILED ({e})")

df = pd.DataFrame(rows, columns=FINAL_COLUMNS)
print(f"Finished: {len(df)} PDB rows | Failed: {len(failed)}")
if failed:
    print("Failed PDB IDs:", failed)


# %% 14. AlphaFold collection
alphafold_rows = extract_alphafold_rows(UNIPROT_ACCESSION, PROTEIN_NAME_COLUMN_VALUE)
print(f"AlphaFold rows: {len(alphafold_rows)}")


# %% 15. Validation checks
def validate_dataframe(df):
    issues = []
    if df.empty and pdb_id_records: issues.append("No PDB rows were produced.")
    if not df.empty:
        if df["PDB ID"].duplicated().any(): issues.append("Duplicate PDB IDs detected.")
        for _, r in df.iterrows():
            if all(r[x] != "N/A" for x in ("Deposited Residue Count", "Modeled Residue Count", "Residue Count Difference")):
                if r["Residue Count Difference"] != r["Deposited Residue Count"] - r["Modeled Residue Count"]:
                    issues.append(f"Residue difference mismatch: {r['PDB ID']}")
    return issues

issues = validate_dataframe(df)
print("Validation:", "PASSED" if not issues else issues)


# %% 16. Excel workbook creation
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, columns):
    for j, col in enumerate(columns, 1):
        c = ws.cell(1, j, col); c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
    ws.row_dimensions[1].height = 28


def autofit(ws, max_width=55):
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, ws.max_row + 1)]
        width = min(max(max(map(len, vals), default=10) + 2, 10), max_width)
        ws.column_dimensions[get_column_letter(col)].width = width

wb = Workbook()
ws = wb.active; ws.title = "PDB_Data"
style_header(ws, FINAL_COLUMNS)

for record in df.to_dict("records"):
    ws.append([record[c] for c in FINAL_COLUMNS])

for row in ws.iter_rows():
    for cell in row:
        cell.border = BORDER; cell.alignment = Alignment(vertical="top", wrap_text=True)

for r in range(2, ws.max_row + 1):
    c = ws.cell(r, FINAL_COLUMNS.index("PDB Link") + 1)
    if c.value:
        pid = str(c.value); c.hyperlink = f"https://www.rcsb.org/structure/{pid}"; c.font = LINK_FONT

ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
autofit(ws)

# AlphaFold sheet
if alphafold_rows or not pdb_id_records:
    af = wb.create_sheet("AlphaFold_Data")
    af_cols = ["Identifier", "Chain", "Positions", "Oligomeric State"]
    style_header(af, af_cols)
    if alphafold_rows:
        for x in alphafold_rows: af.append([x.get(c, "") for c in af_cols])
        for r in range(2, af.max_row + 1):
            c = af.cell(r, 1); ident = str(c.value or "").strip()
            if ident:
                c.hyperlink = ALPHAFOLD_ENTRY_URL.format(identifier=ident); c.font = LINK_FONT
    else:
        af.merge_cells("A3:E4"); af["A3"] = f"No crystal or AlphaFold structure available for {PROTEIN_NAME_COLUMN_VALUE}"; af["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in af.iter_rows():
        for cell in row: cell.border = BORDER; cell.alignment = Alignment(vertical="top", wrap_text=True)
    af.freeze_panes = "A2"; af.auto_filter.ref = af.dimensions if af.max_row > 1 else "A1:E1"; af.sheet_view.showGridLines = False; autofit(af, 45)

# Sources sheet
src = wb.create_sheet("Sources")
src.append(["Category", "Detail"]); style_header(src, ["Category", "Detail"])
source_rows = [
    ("Protein searched", protein_name), ("UniProt accession", UNIPROT_ACCESSION),
    ("UniProt API", UNIPROT_ENTRY_URL.format(accession=UNIPROT_ACCESSION)),
    ("RCSB Search API", RCSB_SEARCH_URL), ("RCSB Data API", "https://data.rcsb.org/"),
    ("wwPDB validation XML", "https://files.rcsb.org/pub/pdb/validation_reports/"),
    ("wwPDB validation PDF", "https://files.rcsb.org/pub/pdb/validation_reports/"),
    ("AlphaFold API", ALPHAFOLD_API_URL.format(accession=UNIPROT_ACCESSION)),
    ("3D-Beacons API", THREED_BEACONS_URL.format(accession=UNIPROT_ACCESSION)),
    ("Retrieval time UTC", datetime.now(timezone.utc).isoformat()),
    ("PDB structures found", len(pdb_id_records)), ("AlphaFold rows", len(alphafold_rows)),
    ("Residue difference", "Deposited polymer monomers minus the sum of modeled_residue_count across all deposited polymer chains."),
    ("R-value rule", "Depositor R-free/R-work from mmCIF; DCC R-free from wwPDB validation XML."),
]
for x in source_rows: src.append(x)
for row in src.iter_rows():
    for cell in row: cell.border = BORDER; cell.alignment = Alignment(vertical="top", wrap_text=True)
src.freeze_panes = "A2"; src.sheet_view.showGridLines = False; src.column_dimensions["A"].width = 32; src.column_dimensions["B"].width = 100

wb.save(OUTPUT_FILENAME)


# %% 17. Final summary
print("\n" + "=" * 70)
print("WORKBOOK CREATED")
print("=" * 70)
print(f"File       : {OUTPUT_FILENAME}")
print(f"UniProt    : {UNIPROT_ACCESSION}")
print(f"PDB rows   : {len(df)}")
print(f"Failed rows: {len(failed)}")
print(f"AlphaFold  : {len(alphafold_rows)}")
print(f"Columns    : {len(FINAL_COLUMNS)}")
print("Notes / Errors column: REMOVED")
print("PDB Chains: next to Structure Type")
print("=" * 70)
