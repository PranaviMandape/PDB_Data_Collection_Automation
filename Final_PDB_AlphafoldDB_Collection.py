# ============================================================
# PDB + AlphaFold Protein Structure Data Collector
# ============================================================
# USER INPUT: change only protein_name.
# Output: Protein_Structure_Data_<protein>.xlsx

# %% 1. Imports and setup
import sys, subprocess, io, gzip, re, time, warnings
import json
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_HALF_EVEN
from collections import defaultdict
from datetime import datetime, timezone
from functools import lru_cache


def ensure_package(import_name, package_name=None):
    try:
        __import__(import_name)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", package_name or import_name])

for mod, pkg in [("requests", "requests"), ("pandas", "pandas"), ("openpyxl", "openpyxl"),
                 ("pdfplumber", "pdfplumber"), ("Bio", "biopython")]:
    ensure_package(mod, pkg)

import requests
import pandas as pd
import pdfplumber
import xml.etree.ElementTree as ET
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from Bio.PDB.MMCIF2Dict import MMCIF2Dict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# %% 2. User input and settings
protein_name = "ITK"                 # <-- CHANGE ONLY THIS
HUMAN_TAXONOMY_ID = 9606
REQUEST_TIMEOUT = 30
MAX_RETRIES = 3
RETRY_BACKOFF = 1.5
RATE_LIMIT_DELAY = 0.15
OUTPUT_FILENAME = f"Protein_Structure_Data_{protein_name}.xlsx"

print(f"Protein: {protein_name}")
print(f"Output:  {OUTPUT_FILENAME}")


# %% 3. API endpoints and HTTP helpers
UNIPROT_SEARCH_URL = "https://rest.uniprot.org/uniprotkb/search"
UNIPROT_ENTRY_URL = "https://rest.uniprot.org/uniprotkb/{accession}.json"
UNIPROT_BROWSER_URL = "https://www.uniprot.org/uniprotkb/{accession}"

RCSB_SEARCH_URL = "https://search.rcsb.org/rcsbsearch/v2/query"
RCSB_ENTRY_URL = "https://data.rcsb.org/rest/v1/core/entry/{pdb_id}"
RCSB_POLYMER_URL = "https://data.rcsb.org/rest/v1/core/polymer_entity/{pdb_id}/{entity_id}"
RCSB_POLYMER_INSTANCE_URL = "https://data.rcsb.org/rest/v1/core/polymer_entity_instance/{pdb_id}/{asym_id}"
RCSB_NONPOLY_URL = "https://data.rcsb.org/rest/v1/core/nonpolymer_entity/{pdb_id}/{entity_id}"
RCSB_ASSEMBLY_URL = "https://data.rcsb.org/rest/v1/core/assembly/{pdb_id}/{assembly_id}"
MMCIF_URL = "https://files.rcsb.org/download/{pdb_id}.cif"

VALIDATION_XML_URL = "https://files.rcsb.org/pub/pdb/validation_reports/{mid}/{pid}/{pid}_validation.xml.gz"
VALIDATION_PDF_URL = "https://files.rcsb.org/pub/pdb/validation_reports/{mid}/{pid}/{pid}_full_validation.pdf.gz"

ALPHAFOLD_API_URL = "https://alphafold.com/api/prediction/{accession}"
ALPHAFOLD_ENTRY_URL = "https://alphafold.com/entry/{identifier}"
THREED_BEACONS_URL = "https://www.ebi.ac.uk/pdbe/pdbe-kb/3dbeacons/api/uniprot/summary/{accession}.json"


def make_session():
    s = requests.Session()
    retry = Retry(total=MAX_RETRIES, connect=MAX_RETRIES, read=MAX_RETRIES,
                  backoff_factor=RETRY_BACKOFF,
                  status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=frozenset({"GET", "POST"}),
                  raise_on_status=False)
    adapter = HTTPAdapter(max_retries=retry)
    s.mount("https://", adapter)
    s.headers.update({"User-Agent": "ProteinPDBDataCollector/2.0"})
    return s

SESSION = make_session()


def request_json(url, params=None, timeout=REQUEST_TIMEOUT):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.get(url, params=params, timeout=timeout)
        r.raise_for_status()
        return r.json(), None
    except (requests.RequestException, ValueError) as e:
        return None, str(e)


def request_bytes(url, timeout=REQUEST_TIMEOUT):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.get(url, timeout=timeout)
        r.raise_for_status()
        return r.content, None
    except requests.RequestException as e:
        return None, str(e)


def rcsb_search(query):
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.post(RCSB_SEARCH_URL, json=query, timeout=REQUEST_TIMEOUT,
                         headers={"Content-Type": "application/json"})
        r.raise_for_status()
        return r.json(), None
    except (requests.RequestException, ValueError) as e:
        return None, str(e)

print("API/session setup complete.")


# %% 4. UniProt human entry resolution
def search_uniprot(gene, reviewed=True):
    q = [f"gene_exact:{gene}", f"organism_id:{HUMAN_TAXONOMY_ID}"]
    if reviewed:
        q.append("reviewed:true")
    data, err = request_json(UNIPROT_SEARCH_URL, {
        "query": " AND ".join(q), "format": "json", "size": 25,
        "fields": "accession,id,protein_name,organism_name,gene_names,reviewed,organism_id"
    })
    return (data or {}).get("results", []) if not err else []


def resolve_uniprot(gene):
    matches = search_uniprot(gene, True) or search_uniprot(gene, False)
    human = [x for x in matches if str(x.get("organism", {}).get("taxonId", x.get("organism_id", ""))) == str(HUMAN_TAXONOMY_ID)]
    if not human:
        raise RuntimeError(f"No human UniProt entry found for exact gene '{gene}'.")
    if len(human) > 1:
        ids = [x.get("primaryAccession", "?") for x in human]
        raise RuntimeError(f"Multiple human UniProt entries found for '{gene}': {ids}")
    x = human[0]
    acc = x.get("primaryAccession")
    if not acc:
        raise RuntimeError("UniProt response did not contain an accession.")
    return acc, HUMAN_TAXONOMY_ID

UNIPROT_ACCESSION, ORGANISM_ID = resolve_uniprot(protein_name)
PROTEIN_NAME_COLUMN_VALUE = f"{protein_name} ({UNIPROT_ACCESSION})"
print(f"UniProt: {UNIPROT_ACCESSION} | Taxonomy: {ORGANISM_ID}")


# %% 5. Exact/display-value helpers
# IMPORTANT:
# Values that are displayed on PDB/wwPDB pages are deliberately stored as TEXT
# in Excel.  This prevents Excel/pandas from converting 2.10 -> 2.1 or 1 -> 1.0.
# The formatting below follows the precision normally displayed for the
# corresponding PDB/wwPDB metric:
#   Resolution                 -> 2 decimal places
#   R-value Free / Work       -> 3 decimal places
#   wwPDB DCC R-free          -> 3 decimal places
#   Clashscore                 -> nearest whole number
#   Rama/sidechain/RSRZ       -> 1 decimal place
#   Average B                  -> value extracted directly from validation PDF
#
# No missing value is replaced with a made-up number; missing values are N/A.

# ROUNDING RULE (this was the cause of the wrong values):
# wwPDB/RCSB store e.g. Ramachandran = 5.65 and Sidechain = 13.85 (2 decimals) and DISPLAY them
# with 1 decimal using "round half to even" -> 5.6 and 13.8.
# The old script used ROUND_HALF_UP -> 5.7 and 13.9 (wrong).
# All arithmetic is done with Decimal on the ORIGINAL TEXT, never through float.
DISPLAY_ROUNDING = ROUND_HALF_EVEN

MISSING_MARKERS = {"", "n/a", "na", ".", "?", "none", "null", "-", "--"}


def raw_value(value):
    """Return a source value as TEXT, using N/A for missing values."""
    if value is None:
        return "N/A"
    value = str(value).strip()
    return "N/A" if value.lower() in MISSING_MARKERS else value


def decimal_from_text(value):
    if value is None or str(value).strip().lower() in MISSING_MARKERS:
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        return None


def display_decimal(value, places=None, integer=False):
    """Format a numeric source value for the value shown by the PDB/wwPDB UI.

    The returned value is ALWAYS a string, so Excel cannot change its
    representation.  Trailing zeroes are retained only when the UI metric
    convention requires fixed precision (e.g. resolution = 2.10).
    """
    d = decimal_from_text(value)
    if d is None:
        return "N/A"

    if integer:
        q = Decimal("1")
        return format(d.quantize(q, rounding=DISPLAY_ROUNDING), "f")

    if places is None:
        return format(d, "f")

    q = Decimal("1").scaleb(-places)
    return format(d.quantize(q, rounding=DISPLAY_ROUNDING), f".{places}f")


def percentage_display(value):
    """PDB validation percentages: one decimal, except zero is shown as 0."""
    d = decimal_from_text(value)
    if d is None:
        return "N/A"
    if d == 0:
        return "0"
    return display_decimal(d, places=1)


def exact_difference(a, b):
    """Calculate a-b using Decimal without binary floating-point rounding."""
    x, y = decimal_from_text(a), decimal_from_text(b)
    if x is None or y is None:
        return "N/A"
    return format(x - y, "f")


# %% 6. RCSB X-ray structure search
# Uses POST because RCSB Search API expects the JSON query as the request body.
def get_pdb_ids(accession):
    query = {
        "query": {
            "type": "group", "logical_operator": "and", "nodes": [
                {"type": "terminal", "service": "text", "parameters": {
                    "attribute": "rcsb_polymer_entity_container_identifiers.reference_sequence_identifiers.database_accession",
                    "operator": "exact_match", "value": accession}},
                {"type": "terminal", "service": "text", "parameters": {
                    "attribute": "exptl.method", "operator": "exact_match", "value": "X-RAY DIFFRACTION"}}
            ]
        },
        "return_type": "entry",
        "request_options": {"paginate": {"start": 0, "rows": 10000}}
    }
    data, err = rcsb_search(query)
    if err:
        raise RuntimeError(f"RCSB Search API failed: {err}")
    ids = []
    for item in (data or {}).get("result_set", []):
        pid = str(item.get("identifier", "")).upper().strip()
        if pid and pid not in ids:
            ids.append(pid)
    return [{"pdb_id": x} for x in ids]

pdb_id_records = get_pdb_ids(UNIPROT_ACCESSION)
print(f"X-ray PDB structures found: {len(pdb_id_records)}")


# %% 7. RCSB entry, method, resolution, residue counts and chains
def get_entry(pdb_id):
    return request_json(RCSB_ENTRY_URL.format(pdb_id=pdb_id))[0]


def get_method(entry):
    methods = [x.get("method", "").strip() for x in entry.get("exptl", []) if x.get("method")]
    return ", ".join(dict.fromkeys(methods)) or "N/A"


def get_resolution(entry):
    vals = (entry.get("rcsb_entry_info") or {}).get("resolution_combined") or []
    vals = vals if isinstance(vals, list) else [vals]
    for value in vals:
        if decimal_from_text(value) is not None:
            # PDB displays resolution to two decimal places.
            return display_decimal(value, places=2)
    return "N/A"


def get_residue_counts(entry, pdb_id):
    """
    Deposited Residue Count = RCSB entry deposited polymer monomer count.
    Modeled Residue Count = SUM of modeled_residue_count for every deposited
    polymer chain (polymer entity instance).
    Difference = Deposited - Modeled.
    """
    deposited_raw = (entry.get("rcsb_entry_info") or {}).get(
        "deposited_polymer_monomer_count"
    )

    try:
        deposited = int(deposited_raw) if deposited_raw is not None else None
    except (TypeError, ValueError):
        deposited = None

    if deposited is None:
        return None, None, None

    polymer_entity_ids = (
        entry.get("rcsb_entry_container_identifiers") or {}
    ).get("polymer_entity_ids") or []

    modeled_total = 0
    found_modeled_value = False
    seen_asym_ids = set()

    for entity_id in polymer_entity_ids:
        entity, err = request_json(
            RCSB_POLYMER_URL.format(
                pdb_id=pdb_id.upper(), entity_id=entity_id
            )
        )
        if err or not entity:
            continue

        asym_ids = (
            entity.get("rcsb_polymer_entity_container_identifiers") or {}
        ).get("asym_ids") or []
        if isinstance(asym_ids, str):
            asym_ids = [asym_ids]

        for asym_id in asym_ids:
            asym_id = str(asym_id).strip()
            if not asym_id or asym_id in seen_asym_ids:
                continue
            seen_asym_ids.add(asym_id)

            instance, err = request_json(
                RCSB_POLYMER_INSTANCE_URL.format(
                    pdb_id=pdb_id.upper(), asym_id=asym_id
                )
            )
            if err or not instance:
                continue

            modeled = (
                instance.get("rcsb_polymer_instance_info") or {}
            ).get("modeled_residue_count")

            if modeled is None:
                continue

            try:
                modeled_total += int(modeled)
                found_modeled_value = True
            except (TypeError, ValueError):
                continue

    modeled = modeled_total if found_modeled_value else None
    difference = deposited - modeled if modeled is not None else None
    return deposited, modeled, difference


def get_pdb_chains(entry, pdb_id):
    ids = (entry.get("rcsb_entry_container_identifiers") or {}).get("polymer_entity_ids") or []
    chains = []
    for entity_id in ids:
        data, err = request_json(RCSB_POLYMER_URL.format(pdb_id=pdb_id.upper(), entity_id=entity_id))
        if err or not data:
            continue
        asym = (data.get("rcsb_polymer_entity_container_identifiers") or {}).get("asym_ids") or []
        asym = [asym] if isinstance(asym, str) else asym
        for chain in asym:
            chain = str(chain).strip()
            if chain and chain not in chains:
                chains.append(chain)
    return ", ".join(chains) if chains else "N/A"


# %% 8. Depositor R-values from mmCIF
def get_depositor_r_values(pdb_id):
    """Read depositor R-values and format them to the PDB page precision."""
    raw, err = request_bytes(MMCIF_URL.format(pdb_id=pdb_id.upper()))
    if err or raw is None:
        return None, None
    try:
        cif = MMCIF2Dict(io.StringIO(raw.decode("utf-8", errors="replace")))

        def first(key):
            v = cif.get(key)
            if isinstance(v, list):
                v = next((x for x in v if str(x).strip() not in MISSING_MARKERS), None)
            if v is None or str(v).strip().lower() in MISSING_MARKERS:
                return None
            return str(v).strip()

        rfree = first("_refine.ls_R_factor_R_free")
        rwork = first("_refine.ls_R_factor_R_work")

        # PDB entry pages display these values to three decimal places.
        return (
            display_decimal(rfree, places=3) if rfree is not None else None,
            display_decimal(rwork, places=3) if rwork is not None else None,
        )
    except Exception:
        return None, None


# %% 9. wwPDB validation XML + Average B-factor
AVERAGE_B_PATTERN = re.compile(r"Average\s*B,\s*all\s*atoms\s*\([^)]*\)\s*(-?\d+\.\d+)", re.I)


def get_validation(pdb_id):
    pid = pdb_id.lower(); mid = pid[1:3]
    raw, err = request_bytes(VALIDATION_XML_URL.format(mid=mid, pid=pid))
    if err or raw is None:
        return {}, None
    try:
        root = ET.fromstring(gzip.decompress(raw))
        node = root.find("Entry")
        return (dict(node.attrib), None) if node is not None else ({}, "no Entry element")
    except Exception as e:
        return {}, str(e)


def get_average_b(pdb_id):
    """Extract Average B, all atoms directly from the wwPDB validation PDF.

    This is intentionally returned as TEXT.  We do not round it in Python
    because the PDF is the source used for the graph/table value.
    """
    pid = pdb_id.lower(); mid = pid[1:3]
    raw, err = request_bytes(VALIDATION_PDF_URL.format(mid=mid, pid=pid))
    if err or raw is None:
        return "N/A"
    try:
        with pdfplumber.open(io.BytesIO(gzip.decompress(raw))) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""

                m = AVERAGE_B_PATTERN.search(text)
                if m:
                    return raw_value(m.group(1))

                for line in text.splitlines():
                    if "Average B" in line and "atom" in line.lower():
                        nums = re.findall(r"-?\d+(?:\.\d+)?", line)
                        if nums:
                            return raw_value(nums[-1])
    except Exception:
        pass
    return "N/A"


# %% 9b. RCSB API validation values (kept as exact TEXT) - used ONLY as a cross-check / fallback
def request_json_exact(url):
    """Like request_json, but numbers are kept as the exact text sent by the server
    (no float conversion, so 5.60 stays '5.60' and nothing is re-rounded)."""
    time.sleep(RATE_LIMIT_DELAY)
    try:
        r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        return json.loads(r.text, parse_float=str, parse_int=str), None
    except (requests.RequestException, ValueError) as e:
        return None, str(e)


def _first_dict(x):
    if isinstance(x, list):
        x = x[0] if x else {}
    return x if isinstance(x, dict) else {}


def get_api_validation_raw(pdb_id):
    """Raw RCSB-API validation numbers (text) for one entry."""
    data, err = request_json_exact(RCSB_ENTRY_URL.format(pdb_id=pdb_id.upper()))
    out = {}
    if err or not data:
        return out
    summ = _first_dict(data.get("pdbx_vrpt_summary"))
    diff = _first_dict(data.get("pdbx_vrpt_summary_diffrn"))
    out["clashscore"] = summ.get("clashscore")
    out["percent-rama-outliers"] = summ.get("percent_ramachandran_outliers")
    out["percent-rota-outliers"] = summ.get("percent_rotamer_outliers")
    out["percent-RSRZ-outliers"] = summ.get("percent_rsrzoutliers")
    out["DCC_Rfree"] = diff.get("dcc_rfree") or summ.get("dcc_rfree")
    return out


# %% 10. Ligand and binding-activity extraction
COMMON_NON_LIGAND_COMPONENTS = {
    "HOH","H2O","WAT","NA","CL","K","MG","CA","ZN","MN","CD","CO","NI","CU","FE","FE2",
    "BR","IOD","CS","LI","SR","BA","RB","AL","AG","AU","HG","SO4","PO4","GOL","EDO","PEG",
    "PG4","MPD","ACT","TRS","IPA","DMS","BME","MES","HEPES","CIT","FMT","ACY","1PE","PGE","P6G",
    "EPE","IMD","NH4","OXL","UNX","UNL","DTT","TAM","BOG","CO3","NO3","AZI"
}


def get_ligand_info(entry, pdb_id):
    ids = (entry.get("rcsb_entry_container_identifiers") or {}).get("non_polymer_entity_ids") or []
    raw = []
    for entity_id in ids:
        data, err = request_json(RCSB_NONPOLY_URL.format(pdb_id=pdb_id, entity_id=entity_id))
        if err or not data:
            continue
        info = data.get("pdbx_entity_nonpoly") or {}
        comp = str(info.get("comp_id", "")).strip()
        if comp:
            raw.append({"comp_id": comp, "name": info.get("name") or comp})
    relevant = [x for x in raw if x["comp_id"].upper() not in COMMON_NON_LIGAND_COMPONENTS]
    return relevant, raw


def get_ligand_activity(entry, ligand_ids):
    affinities = entry.get("rcsb_binding_affinity") or []
    if not affinities or not ligand_ids:
        return "N/A"
    grouped = defaultdict(lambda: defaultdict(list))
    for a in affinities:
        comp = a.get("comp_id", "")
        if comp in ligand_ids:
            grouped[comp][a.get("type", "?")].append((a.get("value"), a.get("unit", "")))
    parts = []
    for comp, types in grouped.items():
        assay_parts = []
        for assay, vals in types.items():
            nums = [v for v, _ in vals if v is not None]
            unit = vals[0][1] if vals else ""
            if len(nums) == 1:
                assay_parts.append(f"{assay}: {nums[0]} ({unit}) from {len(vals)} assay(s)")
            elif nums:
                assay_parts.append(f"{assay}: min: {min(nums)}, max: {max(nums)} ({unit}) from {len(vals)} assay(s)")
        if assay_parts:
            parts.append(f"[{comp}] " + "; ".join(assay_parts))
    return " | ".join(parts) if parts else "N/A"


# %% 11. AlphaFold DB collection

THREED_BEACONS_URL = "https://www.ebi.ac.uk/pdbe/pdbe-kb/3dbeacons/api/uniprot/summary/{accession}.json"
ALPHAFOLD_ENTRY_URL = "https://alphafold.ebi.ac.uk/entry/{identifier}"

def get_alphafold_db_summary(accession):
    data, err = request_json(THREED_BEACONS_URL.format(accession=accession))
    if err or data is None:
        print(f"AlphaFold DB request failed for {accession}: {err}"); return []
    structures = data.get("structures", [])
    if not isinstance(structures, list): return []
    out = []
    for structure in structures:
        if not isinstance(structure, dict): continue
        summary = structure.get("summary", {})
        if not isinstance(summary, dict): continue
        provider = str(summary.get("provider", "")).strip()
        identifier = str(summary.get("model_identifier", "")).strip()
        if provider.lower() != "alphafold db": continue
        if not identifier or not identifier.upper().startswith("AF-"): continue
        out.append(summary)
    return out

def get_af_chain(summary):
    chains = []
    for entity in summary.get("entities", []):
        if not isinstance(entity, dict): continue
        chain_ids = entity.get("chain_ids", [])
        if isinstance(chain_ids, str): chain_ids = [chain_ids]
        if not isinstance(chain_ids, list): continue
        for chain in chain_ids:
            chain = str(chain).strip()
            if chain and chain not in {".", "?"} and chain not in chains: chains.append(chain)
    return ", ".join(chains) if chains else "N/A"

def get_af_positions(summary):
    start = summary.get("uniprot_start")
    end = summary.get("uniprot_end")
    if start is None or end is None: return "N/A"
    try: return f"{int(start)}-{int(end)}"
    except (TypeError, ValueError): return "N/A"

def get_af_oligomeric_state(summary):
    value = summary.get("oligomeric_state")
    if value is None: return "N/A"
    value = str(value).strip()
    return value if value else "N/A"

def extract_alphafold_rows(accession, protein_display):
    structures = get_alphafold_db_summary(accession)
    rows, seen = [], set()
    for summary in structures:
        identifier = str(summary.get("model_identifier", "")).strip()
        if not identifier: continue
        key = identifier.upper()
        if key in seen: continue
        seen.add(key)
        rows.append({
            "Protein Name": protein_display,
            "Identifier": identifier,
            "Chain": get_af_chain(summary),
            "Positions": get_af_positions(summary),
            "Oligomeric State": get_af_oligomeric_state(summary)
        })
    rows.sort(key=lambda x: str(x.get("Identifier", "")).upper())
    return rows


# %% 12. Column definitions (Notes / Errors removed)
FINAL_COLUMNS = [
    "Protein Name", "PDB ID", "PDB Link", "Method", "Resolution (Å)",
    "Deposited Residue Count", "Modeled Residue Count", "Residue Count Difference",
    "R-Value Free", "R-Value Work", "R-Value Difference", "R-free",
    "Clashscore", "Ramachandran Outliers (%)", "Sidechain Outliers (%)", "RSRZ Outliers (%)",
    "Average B, all atoms (Å²)", "Ligand ID", "Ligand Name", "Structure Type", "PDB Chains", "Ligand Activity",
    "All HET Components (raw)"
]

print(f"PDB_Data columns: {len(FINAL_COLUMNS)}")


# %% 13. Main PDB processing
rows = []
failed = []
audit_rows = []
for i, rec in enumerate(pdb_id_records, 1):
    pid = rec["pdb_id"]
    print(f"[{i}/{len(pdb_id_records)}] {pid} ...", end=" ")
    row = {c: "N/A" for c in FINAL_COLUMNS}
    row["Protein Name"] = PROTEIN_NAME_COLUMN_VALUE
    row["PDB ID"] = pid
    row["PDB Link"] = pid
    try:
        entry, err = request_json(RCSB_ENTRY_URL.format(pdb_id=pid))
        if err or not entry:
            failed.append((pid, f"RCSB entry: {err}")); rows.append(row); print("FAILED"); continue

        row["Method"] = get_method(entry)
        row["Resolution (Å)"] = get_resolution(entry)
        dep, mod, diff = get_residue_counts(entry, pid)
        row["Deposited Residue Count"] = dep if dep is not None else "N/A"
        row["Modeled Residue Count"] = mod if mod is not None else "N/A"
        row["Residue Count Difference"] = diff if diff is not None else "N/A"

        chains = get_pdb_chains(entry, pid)
        validation, _ = get_validation(pid)
        rfree, rwork = get_depositor_r_values(pid)
        row["R-Value Free"] = raw_value(rfree)
        row["R-Value Work"] = raw_value(rwork)
        row["R-Value Difference"] = exact_difference(
            row["R-Value Free"], row["R-Value Work"]
        )

        # wwPDB validation metrics: preserve the page-display precision.
        row["R-free"] = display_decimal(validation.get("DCC_Rfree"), places=3)
        row["Clashscore"] = display_decimal(
            validation.get("clashscore"), integer=True
        )
        row["Ramachandran Outliers (%)"] = percentage_display(
            validation.get("percent-rama-outliers")
        )
        row["Sidechain Outliers (%)"] = percentage_display(
            validation.get("percent-rota-outliers")
        )
        row["RSRZ Outliers (%)"] = percentage_display(
            validation.get("percent-RSRZ-outliers")
        )
        row["Average B, all atoms (Å²)"] = raw_value(get_average_b(pid))

        # ---- audit trail: raw source text vs. final value written to Excel
        api_raw = get_api_validation_raw(pid)
        for metric, col in (("DCC_Rfree", "R-free"), ("clashscore", "Clashscore"),
                            ("percent-rama-outliers", "Ramachandran Outliers (%)"),
                            ("percent-rota-outliers", "Sidechain Outliers (%)"),
                            ("percent-RSRZ-outliers", "RSRZ Outliers (%)")):
            xml_txt = validation.get(metric)
            api_txt = api_raw.get(metric)
            api_fmt = (display_decimal(api_txt, places=3) if metric == "DCC_Rfree" else
                       display_decimal(api_txt, integer=True) if metric == "clashscore" else
                       percentage_display(api_txt)) if api_txt is not None else "N/A"
            audit_rows.append({"PDB ID": pid, "Column": col,
                               "Raw value (validation XML)": raw_value(xml_txt),
                               "Raw value (RCSB API)": raw_value(api_txt),
                               "Value written to sheet": row[col],
                               "API-formatted value": api_fmt,
                               "Sources agree?": "YES" if api_fmt == row[col] else ("n/a" if api_fmt == "N/A" else "CHECK")})
        # Fallback: if the XML had nothing for a metric but the API does, use the API text.
        for metric, col in (("DCC_Rfree", "R-free"), ("clashscore", "Clashscore"),
                            ("percent-rama-outliers", "Ramachandran Outliers (%)"),
                            ("percent-rota-outliers", "Sidechain Outliers (%)"),
                            ("percent-RSRZ-outliers", "RSRZ Outliers (%)")):
            if row[col] == "N/A" and api_raw.get(metric) is not None:
                t = api_raw[metric]
                row[col] = (display_decimal(t, places=3) if metric == "DCC_Rfree" else
                            display_decimal(t, integer=True) if metric == "clashscore" else
                            percentage_display(t))

        ligands, raw = get_ligand_info(entry, pid)
        row["All HET Components (raw)"] = ", ".join(x["comp_id"] for x in raw) if raw else "N/A"
        if ligands:
            row["Ligand ID"] = ", ".join(x["comp_id"] for x in ligands)
            row["Ligand Name"] = ", ".join(x["name"] for x in ligands)
            row["Structure Type"] = "Complex"
            row["Ligand Activity"] = get_ligand_activity(entry, {x["comp_id"] for x in ligands})
        else:
            row["Structure Type"] = "Apo"
        # PDB Chains intentionally placed next to Structure Type in the final Excel order.
        row["PDB Chains"] = chains
        # Force all display-sensitive fields to TEXT before DataFrame creation.
        for _col in (
            "Resolution (Å)", "R-Value Free", "R-Value Work", "R-Value Difference",
            "R-free", "Clashscore", "Ramachandran Outliers (%)",
            "Sidechain Outliers (%)", "RSRZ Outliers (%)",
            "Average B, all atoms (Å²)"
        ):
            row[_col] = raw_value(row.get(_col))

        rows.append(row)
        print("done")
    except Exception as e:
        failed.append((pid, str(e))); rows.append(row); print(f"FAILED ({e})")

df = pd.DataFrame(rows, columns=FINAL_COLUMNS)
print(f"Finished: {len(df)} PDB rows | Failed: {len(failed)}")
if failed:
    print("Failed PDB IDs:", failed)


# %% 14. AlphaFold collection
alphafold_rows = extract_alphafold_rows(UNIPROT_ACCESSION, PROTEIN_NAME_COLUMN_VALUE)
print(f"AlphaFold rows: {len(alphafold_rows)}")


# %% 15. Validation checks
def validate_dataframe(df):
    issues = []
    if df.empty and pdb_id_records: issues.append("No PDB rows were produced.")
    if not df.empty:
        if df["PDB ID"].duplicated().any(): issues.append("Duplicate PDB IDs detected.")
        for _, r in df.iterrows():
            if all(r[x] != "N/A" for x in ("Deposited Residue Count", "Modeled Residue Count", "Residue Count Difference")):
                if r["Residue Count Difference"] != r["Deposited Residue Count"] - r["Modeled Residue Count"]:
                    issues.append(f"Residue difference mismatch: {r['PDB ID']}")
    return issues

issues = validate_dataframe(df)
print("Validation:", "PASSED" if not issues else issues)


# %% 16. Excel workbook creation
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, columns):
    for j, col in enumerate(columns, 1):
        c = ws.cell(1, j, col); c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
    ws.row_dimensions[1].height = 28


def autofit(ws, max_width=55):
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, ws.max_row + 1)]
        width = min(max(max(map(len, vals), default=10) + 2, 10), max_width)
        ws.column_dimensions[get_column_letter(col)].width = width

wb = Workbook()
ws = wb.active; ws.title = "PDB_Data"
style_header(ws, FINAL_COLUMNS)

DISPLAY_TEXT_COLUMNS = {
    "Resolution (Å)", "R-Value Free", "R-Value Work", "R-Value Difference",
    "R-free", "Clashscore", "Ramachandran Outliers (%)",
    "Sidechain Outliers (%)", "RSRZ Outliers (%)",
    "Average B, all atoms (Å²)"
}

for record in df.to_dict("records"):
    ws.append([record[c] for c in FINAL_COLUMNS])

for row in ws.iter_rows():
    for cell in row:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# Excel number format '@' makes these cells text even when the value looks numeric.
for col_name in DISPLAY_TEXT_COLUMNS:
    col_idx = FINAL_COLUMNS.index(col_name) + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, col_idx).number_format = "@"

for r in range(2, ws.max_row + 1):
    c = ws.cell(r, FINAL_COLUMNS.index("PDB Link") + 1)
    if c.value:
        pid = str(c.value); c.hyperlink = f"https://www.rcsb.org/structure/{pid}"; c.font = LINK_FONT

ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
autofit(ws)

# AlphaFold sheet

af = wb.create_sheet("AlphaFold_Data")
af_cols = ["Protein Name", "Identifier", "Chain", "Positions", "Oligomeric State"]
style_header(af, af_cols)

if alphafold_rows:
    for record in alphafold_rows:
        af.append([record.get(c, "N/A") for c in af_cols])

    identifier_col = af_cols.index("Identifier") + 1
    for r in range(2, af.max_row + 1):
        cell = af.cell(r, identifier_col)
        identifier = str(cell.value or "").strip()
        if identifier and identifier != "N/A":
            cell.hyperlink = ALPHAFOLD_ENTRY_URL.format(identifier=identifier)
            cell.font = LINK_FONT
else:
    af.append([PROTEIN_NAME_COLUMN_VALUE, "N/A", "N/A", "N/A", "N/A"])

for row in af.iter_rows():
    for cell in row:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)

af.freeze_panes = "A2"
af.auto_filter.ref = af.dimensions
af.sheet_view.showGridLines = False
autofit(af, 45)

# Audit sheet: raw source text next to the value written, so any value can be verified
au = wb.create_sheet("Validation_Audit")
au_cols = ["PDB ID", "Column", "Raw value (validation XML)", "Raw value (RCSB API)",
           "Value written to sheet", "API-formatted value", "Sources agree?"]
style_header(au, au_cols)
for rec in audit_rows:
    au.append([rec[c] for c in au_cols])
for row_ in au.iter_rows():
    for cell in row_:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if cell.row > 1:
            cell.number_format = "@"
au.freeze_panes = "A2"; au.auto_filter.ref = au.dimensions; au.sheet_view.showGridLines = False
autofit(au, 35)
n_check = sum(1 for r in audit_rows if r["Sources agree?"] == "CHECK")
print(f"Audit: {n_check} value(s) where XML and API disagree (see Validation_Audit sheet, filter 'CHECK').")
wb.save(OUTPUT_FILENAME)


# %% 17. Final summary
print("\n" + "=" * 70)
print("WORKBOOK CREATED")
print("=" * 70)
print(f"File       : {OUTPUT_FILENAME}")
print(f"UniProt    : {UNIPROT_ACCESSION}")
print(f"PDB rows   : {len(df)}")
print(f"Failed rows: {len(failed)}")
print(f"AlphaFold  : {len(alphafold_rows)}")
print(f"Columns    : {len(FINAL_COLUMNS)}")
print("=" * 70)
